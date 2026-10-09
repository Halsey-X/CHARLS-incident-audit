# -*- coding: utf-8 -*-
"""Recompute Table 3 (multimorbidity rates) from the raw derivation data:
distribution of the number of the 14 chronic conditions per participant."""
import os, json
from collections import Counter
import openpyxl

XLSX = r'D:\02XIHO Files\Daily Work\03Scientific Research\00Article_File\00Student\01单任务论文\多任务原始数据\2011_2013关联规则.xlsx'
OUT = r'D:\02XIHO Files\Daily Work\03Scientific Research\00Article_File\00Student\01单任务论文\PublicHealth_Submission\04_Data_and_Results\table3_recomputed.json'

wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
ws = wb[wb.sheetnames[0]]
it = ws.iter_rows(values_only=True)
hdr = list(next(it))
cols = list(range(len(hdr)))
DISEASE = cols[4:18]       # 14 chronic conditions
WAVE = 33                  # 数据年份

rows = list(it)
print('total rows:', len(rows))
print('wave distribution:', Counter(str(r[WAVE]) for r in rows))
print('disease cols:', [hdr[i] for i in DISEASE])

def yes(v):
    if v is None:
        return False
    s = str(v).strip()
    return s in ('是', '1', '1.0', 'Yes', 'yes', 'TRUE', 'True')

counts = Counter()
missing_informative = 0
for r in rows:
    vals = [r[i] for i in DISEASE]
    # skip rows with no usable value at all
    if all(v is None for v in vals):
        missing_informative += 1
        continue
    k = sum(1 for v in vals if yes(v))
    counts[k] += 1

N = sum(counts.values())
print('informative rows:', N, ' (all-None rows skipped:', missing_informative, ')')

dist = {}
for k in sorted(counts):
    dist[k] = counts[k]
print('raw distribution of disease count:')
for k in sorted(counts):
    print(f'   {k:>2} : {counts[k]:>6}  ({counts[k]/N*100:.2f}%)')

def pct(x):
    return round(x / N * 100, 2)

tab = [
    ('0',            counts.get(0, 0)),
    ('1',            counts.get(1, 0)),
    ('2',            counts.get(2, 0)),
    ('>=3',          sum(v for k, v in counts.items() if k >= 3)),
]
print()
print('=== Table 3 candidate rows ===')
for lab, n in tab:
    print(f'   {lab:>4} | {n:>6} | {pct(n):>6.2f}')
tot2 = counts.get(2, 0) + sum(v for k, v in counts.items() if k >= 3)
print(f'   checked: no-disease {pct(counts.get(0,0))}, >=2 {pct(tot2)}, >=3 {pct(sum(v for k,v in counts.items() if k>=3))}')
print(f'   sum check: {sum(n for _, n in tab)} == {N}')

json.dump({'N': N,
           'dist': {str(k): v for k, v in dist.items()},
           'table3': [{'label': l, 'n': n, 'pct': pct(n)} for l, n in tab],
           'check': {'zero': pct(counts.get(0, 0)), 'ge2': pct(tot2),
                     'ge3': pct(sum(v for k, v in counts.items() if k >= 3))}},
          open(OUT, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
print('saved ->', OUT)
