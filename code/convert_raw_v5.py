# -*- coding: utf-8 -*-
"""Convert the two raw CHARLS panel workbooks to CSV for the v5 re-analysis."""
import openpyxl, csv, os, time

RAW = r"D:\02XIHO Files\Daily Work\03Scientific Research\00Article_File\00Student\01单任务论文\多任务原始数据"
OUT = r"D:\02XIHO Files\Daily Work\03Scientific Research\00Article_File\00Student\01单任务论文\PublicHealth_Submission\04_Data_and_Results\v5"
os.makedirs(OUT, exist_ok=True)

for src, dst in [("2011_2013关联规则.xlsx", "raw_wave_2011_2013_2015.csv"),
                 ("15-18外部验证.xlsx", "raw_wave_2015_2018_2020.csv")]:
    t0 = time.time()
    wb = openpyxl.load_workbook(os.path.join(RAW, src), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    n = 0
    with open(os.path.join(OUT, dst), "w", encoding="utf-8-sig", newline="") as fh:
        w = csv.writer(fh)
        for row in ws.iter_rows(values_only=True):
            w.writerow(["" if v is None else v for v in row])
            n += 1
    wb.close()
    print("%-34s -> %-28s rows=%5d  %.1fs" % (src, dst, n - 1, time.time() - t0))
print("CONVERT_DONE")
